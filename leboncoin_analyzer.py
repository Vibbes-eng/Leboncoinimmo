#!/usr/bin/env python3
"""
LeBonCoin Analyzer — v3
========================
Améliorations v3 :
- Tous les résultats sans filtre (filtrage dans HTML + Excel)
- Export Excel .xlsx : AutoFilter, freeze pane, mise en forme conditionnelle
- Cache permanent des annonces déjà scrapées (skip au relancement)
- DPE extrait et coloré (F/G signalés en rouge — opportunité travaux)
- Taxe foncière mieux extraite + estimation indicative si absente
- Prompt IA anti-hallucination réécrit par prompt engineer
- Tri colonnes HTML avec indicateurs ▲/▼, valeurs vides repoussées en bas
- Loyer minimum requis affiché pour chaque bien sans loyer
"""

import asyncio
import csv
import json
import math
import os
import random
import re
import sys
import urllib.parse
import urllib.request
from datetime import datetime
from pathlib import Path

import anthropic as _anthropic
import openai as _openai
from dotenv import load_dotenv
from playwright.async_api import async_playwright

try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles.differential import DifferentialStyle
    XLSX_OK = True
except ImportError:
    XLSX_OK = False

# ── Configuration ─────────────────────────────────────────────────
load_dotenv()

CDP_PORT      = int(os.getenv("CDP_PORT", "9222"))
MIN_YIELD     = float(os.getenv("MIN_YIELD", "10.0"))   # sert dans le prompt IA, plus de filtre dur
MAX_PRICE     = int(os.getenv("MAX_PRICE", "150000"))
WA_PHONE      = os.getenv("WA_PHONE", "").strip()
WA_APIKEY     = os.getenv("WA_APIKEY", "").strip()
ANTHROPIC_KEY = os.getenv("ANTHROPIC_API_KEY", "").strip()
OPENAI_KEY    = os.getenv("OPENAI_API_KEY", "").strip()

AI_PROVIDER   = "openai" if OPENAI_KEY else ("anthropic" if ANTHROPIC_KEY else None)

CDP_URL       = f"http://localhost:{CDP_PORT}"
SEEN_FILE     = Path("seen_urls.json")
CACHE_FILE    = Path("cache_annonces.json")
COOKIES_FILE  = Path("cookies.json")
OUTPUT_DIR    = Path(".")

BLANK = "—"

DPE_COLORS = {
    "A": "#009966", "B": "#33cc33", "C": "#99cc00",
    "D": "#ffcc00", "E": "#ff9900", "F": "#ff5500", "G": "#cc0000",
}

# ── User-Agents et Viewports ───────────────────────────────────────
_USER_AGENTS = [
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/121.0.6167.184 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.6099.225 Safari/537.36",
]
_VIEWPORTS = [
    {"width": 1920, "height": 1080},
    {"width": 1680, "height": 1050},
    {"width": 1536, "height": 864},
    {"width": 1440, "height": 900},
    {"width": 1366, "height": 768},
]

# ── Stealth JS v2 ─────────────────────────────────────────────────
_STEALTH_JS = """
(function() {
  'use strict';
  Object.defineProperty(navigator, 'webdriver', {get: () => undefined});
  const makePlugin = (name, desc, filename) => {
    const p = Object.create(Plugin.prototype);
    Object.defineProperty(p, 'name', {get: () => name});
    Object.defineProperty(p, 'description', {get: () => desc});
    Object.defineProperty(p, 'filename', {get: () => filename});
    Object.defineProperty(p, 'length', {get: () => 0});
    return p;
  };
  const plugins = [
    makePlugin('Chrome PDF Plugin', 'Portable Document Format', 'internal-pdf-viewer'),
    makePlugin('Chrome PDF Viewer', '', 'mhjfbmdgcfjbbpaeojofohoefgiehjai'),
    makePlugin('Native Client', '', 'internal-nacl-plugin'),
  ];
  Object.defineProperty(navigator, 'plugins', {
    get: () => { const a=[...plugins]; a.__proto__=PluginArray.prototype; return a; }
  });
  Object.defineProperty(navigator, 'languages', {get: () => ['fr-FR','fr','en-US','en']});
  if (!window.chrome) {
    window.chrome = {
      runtime: {connect:()=>{}, sendMessage:()=>{}, onMessage:{addListener:()=>{},removeListener:()=>{}}},
      loadTimes: () => ({requestTime:Date.now()/1000-0.2, wasNpnNegotiated:true, npnNegotiatedProtocol:'h2'}),
      csi: () => ({startE:Date.now(), onloadT:Date.now()+300, pageT:700}),
      app: {}
    };
  }
  try {
    const orig = window.navigator.permissions.query.bind(navigator.permissions);
    window.navigator.permissions.query = (p) =>
      p.name==='notifications' ? Promise.resolve({state:Notification.permission}) : orig(p);
  } catch(e) {}
  ['cdc_adoQpoasnfa76pfcZLmcfl_Array','cdc_adoQpoasnfa76pfcZLmcfl_Promise',
   'cdc_adoQpoasnfa76pfcZLmcfl_Symbol','__playwright_target__','__pw_manual'].forEach(k=>{
    try{delete window[k];}catch(e){}
  });
  const origToDataURL = HTMLCanvasElement.prototype.toDataURL;
  HTMLCanvasElement.prototype.toDataURL = function(type) {
    if (type==='image/png' && this.width===1 && this.height===1) return origToDataURL.apply(this,arguments);
    const ctx = this.getContext('2d');
    if (ctx) { const d=ctx.getImageData(0,0,1,1); d.data[0]=Math.max(0,d.data[0]-1); ctx.putImageData(d,0,0); }
    return origToDataURL.apply(this,arguments);
  };
  const gp = WebGLRenderingContext.prototype.getParameter;
  WebGLRenderingContext.prototype.getParameter = function(p) {
    if (p===37445) return 'Intel Inc.';
    if (p===37446) return 'Intel(R) Iris(TM) Plus Graphics 640';
    return gp.apply(this,[p]);
  };
  Object.defineProperty(navigator,'hardwareConcurrency',{get:()=>8});
  Object.defineProperty(navigator,'deviceMemory',{get:()=>8});
  Object.defineProperty(navigator,'platform',{get:()=>'Win32'});
  Object.defineProperty(navigator,'maxTouchPoints',{get:()=>0});
  Object.defineProperty(screen,'colorDepth',{get:()=>24});
  Date.prototype.getTimezoneOffset = function(){return -60;};
  if (navigator.getBattery) {
    navigator.getBattery = () => Promise.resolve({
      charging:true, chargingTime:0, dischargingTime:Infinity, level:0.98,
      addEventListener:()=>{}, removeEventListener:()=>{}
    });
  }
})();
"""

# ── Stealth + headers ──────────────────────────────────────────────
async def apply_stealth(context, ua: str = None, viewport: dict = None) -> None:
    chosen_ua = ua or random.choice(_USER_AGENTS)
    chosen_vp = viewport or random.choice(_VIEWPORTS)
    await context.add_init_script(_STEALTH_JS)
    ver_match = re.search(r"Chrome/(\d+)", chosen_ua)
    chrome_ver = ver_match.group(1) if ver_match else "124"
    await context.set_extra_http_headers({
        "Accept-Language": "fr-FR,fr;q=0.9,en-US;q=0.8,en;q=0.7",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
        "Accept-Encoding": "gzip, deflate, br",
        "sec-ch-ua": f'"Google Chrome";v="{chrome_ver}", "Chromium";v="{chrome_ver}", "Not-A.Brand";v="99"',
        "sec-ch-ua-mobile": "?0",
        "sec-ch-ua-platform": '"Windows"',
        "sec-fetch-dest": "document",
        "sec-fetch-mode": "navigate",
        "sec-fetch-site": "none",
        "sec-fetch-user": "?1",
        "Upgrade-Insecure-Requests": "1",
        "User-Agent": chosen_ua,
        "DNT": "1",
    })
    print(f"  [Stealth] UA: Chrome/{chrome_ver} | Viewport: {chosen_vp['width']}x{chosen_vp['height']}")

# ── Comportement humain ────────────────────────────────────────────
async def human_mouse_move(page, target_x=None, target_y=None):
    vp = page.viewport_size or {"width": 1366, "height": 768}
    w, h = vp["width"], vp["height"]
    sx, sy = random.randint(100, w-100), random.randint(100, h-100)
    ex = target_x or random.randint(200, w-200)
    ey = target_y or random.randint(200, h-200)
    cp1x = sx + (ex-sx)*random.uniform(0.2,0.4) + random.randint(-80,80)
    cp1y = sy + (ey-sy)*random.uniform(0.1,0.3) + random.randint(-80,80)
    cp2x = sx + (ex-sx)*random.uniform(0.6,0.8) + random.randint(-80,80)
    cp2y = sy + (ey-sy)*random.uniform(0.7,0.9) + random.randint(-80,80)
    for i in range(random.randint(15,30)+1):
        t = i / 25
        x = (1-t)**3*sx + 3*(1-t)**2*t*cp1x + 3*(1-t)*t**2*cp2x + t**3*ex
        y = (1-t)**3*sy + 3*(1-t)**2*t*cp1y + 3*(1-t)*t**2*cp2y + t**3*ey
        await page.mouse.move(int(x), int(y))
        await page.wait_for_timeout(random.randint(10, 35))

async def human_scroll(page, direction="down"):
    total = random.randint(300, 800)
    chunks = random.randint(4, 9)
    for _ in range(chunks):
        delta = total // chunks + random.randint(-30, 30)
        await page.mouse.wheel(0, delta if direction == "down" else -delta)
        await page.wait_for_timeout(random.randint(80, 220))

def human_delay(min_ms=800, max_ms=2500):
    return int(min_ms + random.betavariate(2, 3) * (max_ms - min_ms))

# ── Cookies ────────────────────────────────────────────────────────
async def save_cookies(context):
    try:
        cookies = await context.cookies()
        COOKIES_FILE.write_text(json.dumps(cookies, indent=2, ensure_ascii=False), encoding="utf-8")
    except Exception as e:
        print(f"  [Cookies] Sauvegarde échouée : {e}")

async def load_cookies(context):
    if not COOKIES_FILE.exists():
        return
    try:
        cookies = json.loads(COOKIES_FILE.read_text(encoding="utf-8"))
        now = datetime.now().timestamp()
        valid = [c for c in cookies if c.get("expires", 0) in (-1, 0) or c.get("expires", 0) > now]
        if valid:
            await context.add_cookies(valid)
            print(f"  [Cookies] {len(valid)} cookie(s) restauré(s)")
    except Exception as e:
        print(f"  [Cookies] Chargement échoué : {e}")

# ── Cache annonces ─────────────────────────────────────────────────
def load_cache() -> dict:
    if CACHE_FILE.exists():
        try:
            return json.loads(CACHE_FILE.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}

def save_cache(cache: dict):
    CACHE_FILE.write_text(json.dumps(cache, indent=2, ensure_ascii=False), encoding="utf-8")

# ── URLs déjà vues ─────────────────────────────────────────────────
def load_seen() -> set:
    if SEEN_FILE.exists():
        try:
            return set(json.loads(SEEN_FILE.read_text(encoding="utf-8")))
        except Exception:
            pass
    return set()

def save_seen(seen: set):
    SEEN_FILE.write_text(json.dumps(sorted(seen), indent=2, ensure_ascii=False), encoding="utf-8")

# ── Détection blocage ──────────────────────────────────────────────
async def detect_and_handle_block(page) -> bool:
    url = page.url
    title = await page.title()
    content = await page.content()
    signals = [
        "captcha" in url.lower(), "challenge" in url.lower(),
        "just a moment" in title.lower(), "veuillez patienter" in title.lower(),
        "accès refusé" in title.lower(), "access denied" in title.lower(),
        'id="challenge-form"' in content, "cf-browser-verification" in content,
        "__cf_chl" in content, "recaptcha" in content.lower(),
        "cloudflare" in content.lower() and "checking" in title.lower(),
    ]
    if any(signals):
        print("\n  ⚠️  BLOCAGE DÉTECTÉ (captcha/Cloudflare)")
        print("  → Résolvez le captcha manuellement dans Chrome")
        print("  → Appuyez sur ENTRÉE quand c'est fait...")
        input()
        await page.wait_for_load_state("networkidle")
        return True
    return False

# ── Détection type de page ─────────────────────────────────────────
async def detect_page_type(page) -> str:
    url = page.url
    if "mes-favoris/recherches" in url:
        return "saved_searches"
    if "recherche" in url or "ventes_immobilieres" in url or "locations" in url:
        return "results"
    await page.wait_for_load_state("domcontentloaded")
    url = page.url
    if "mes-favoris/recherches" in url:
        return "saved_searches"
    if "recherche" in url or "ventes_immobilieres" in url or "locations" in url:
        return "results"
    return "unknown"

async def get_saved_search_urls(page) -> list:
    await page.wait_for_load_state("networkidle")
    urls = []
    for sel in ["a[href*='recherche']","a[href*='ventes_immobilieres']","a[href*='locations']",
                "[data-qa-id='saved-search-link']","a[data-test-id='saved-search-link']"]:
        for link in await page.query_selector_all(sel):
            href = await link.get_attribute("href")
            if href and any(k in href for k in ["recherche","ventes_immobilieres","locations"]):
                full = href if href.startswith("http") else "https://www.leboncoin.fr" + href
                if full not in urls:
                    urls.append(full)
    return urls

# ── Parsing ────────────────────────────────────────────────────────
def _text(s): return s.strip() if s else ""

def parse_price(text):
    if not text: return None
    c = re.sub(r"[^\d]", "", text)
    return int(c) if c else None

def parse_surface(text):
    if not text: return None
    m = re.search(r"(\d+[\.,]?\d*)\s*m", text)
    return float(m.group(1).replace(",", ".")) if m else None

# ── Prompt IA v3 — anti-hallucination, prompt engineer ────────────
_AI_PROMPT = """\
RÔLE : Tu es un expert en analyse d'annonces immobilières françaises pour investisseurs \
locatifs qui ciblent une rentabilité brute ≥ {min_yield}%.

━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
RÈGLE ABSOLUE — NE JAMAIS HALLUCINER
• Si une donnée n'est PAS écrite noir sur blanc dans le texte → null
• INTERDIT : estimer, calculer, déduire ou inventer un loyer ou un chiffre financier
• Un silence dans le texte = null dans le JSON. Jamais une supposition.
━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━

CONTEXTE DE L'ANNONCE :
• Prix de vente : {price_hint}
• Surface : {surface_hint}
• Localisation : {location_hint}

TEXTE DE L'ANNONCE :
\"\"\"
{description}
\"\"\"

RÈGLES D'EXTRACTION :

[LOYER]
✓ Chercher : "loyer", "loué", "bail", "occupé", "revenus locatifs", "rapport locatif",
  "charges comprises", "hors charges", "CC", "HC", "/mois", "mensualité locative"
✓ Format "740 x 12 = 8 880€" → loyer = 740 €/mois (le x12 est l'annualisation)
✓ CC (charges comprises) → loyer_cc | HC (hors charges) → loyer_hc
✓ Si loyer_cc ET charges connues → loyer_hc = loyer_cc - charges
✓ Immeuble / plusieurs lots → additionner TOUS les loyers → loyer_hc total
✗ JAMAIS confondre prix de vente et loyer
✗ JAMAIS déduire le loyer depuis le prix ou la surface — c'est de l'hallucination
✗ Si "loyer potentiel" ou "estimation loyer" → loyer_potentiel uniquement (pas loyer_hc)

[TAXE FONCIÈRE]
✓ Chercher : "taxe foncière", "TF :", "TF annuelle", "taxe fonc"
✗ Si absente → null (ne pas estimer)

[DPE]
✓ Chercher : "DPE", "diagnostic", "classe énergie", "consommation"
✓ Lettre unique A/B/C/D/E/F/G associée à ces mots
✗ Si non mentionné → null

[STATUT / BAIL]
✓ "loué", "occupé", "locataire en place", "bail en cours" → deja_loue = true
✓ "meublé", "nu", "vide", "commercial", "professionnel" → bail_type correspondant

RETOURNE UNIQUEMENT CE JSON (aucun texte avant ou après) :
{{
  "loyer_hc": <entier €/mois ou null>,
  "loyer_cc": <entier €/mois ou null>,
  "charges_mois": <entier €/mois ou null>,
  "taxe_fonciere": <entier €/an ou null>,
  "loyer_potentiel": <entier si "estimé"/"potentiel" mentionné, sinon null>,
  "nb_lots": <entier — nombre de logements, 1 si seul>,
  "loyer_detail": <liste d'entiers si plusieurs lots, sinon null>,
  "loyer_source": "<hc|cc|cc_moins_charges|calcule|potentiel|inconnu>",
  "deja_loue": <true|false>,
  "bail_type": "<vide|meuble|commercial|inconnu>",
  "dpe": "<A|B|C|D|E|F|G|null>",
  "dpe_valeur_kwh": <entier kWh/m²/an ou null>,
  "notes": "<1 phrase max — ce qui a été trouvé ou pourquoi null>"
}}"""

# ── Extraction IA ──────────────────────────────────────────────────
async def ai_extract_financials(description: str, price=None, surface=None, location=None) -> dict:
    if not AI_PROVIDER or not description or description == BLANK:
        return {}
    price_hint    = f"{price:,}€".replace(",", "\u202f") if price else "non précisé"
    surface_hint  = f"{surface} m²" if surface else "non précisée"
    location_hint = location or "non précisée"
    prompt = _AI_PROMPT.format(
        min_yield=MIN_YIELD, price_hint=price_hint,
        surface_hint=surface_hint, location_hint=location_hint,
        description=description[:3000],
    )
    try:
        if AI_PROVIDER == "openai":
            client = _openai.AsyncOpenAI(api_key=OPENAI_KEY)
            resp = await client.chat.completions.create(
                model="gpt-4o-mini", max_tokens=600, temperature=0,
                messages=[{"role": "user", "content": prompt}],
            )
            raw = resp.choices[0].message.content.strip()
        else:
            client = _anthropic.AsyncAnthropic(api_key=ANTHROPIC_KEY)
            resp = await client.messages.create(
                model="claude-haiku-4-5-20251001", max_tokens=600,
                messages=[{"role": "user", "content": prompt}],
            )
            raw = resp.content[0].text.strip()
        m = re.search(r"\{.*\}", raw, re.DOTALL)
        if m:
            return json.loads(m.group())
    except Exception as e:
        print(f"    [AI/{AI_PROVIDER}] Erreur : {e}")
    return {}

# ── Fallbacks regex ────────────────────────────────────────────────
def _regex_fallback_loyer(text):
    if not text: return None
    for pat in [
        r"loyer\s+hors\s+charges[^\d]*(\d[\d\s]{2,5})",
        r"loyer\s+hc[^\d]*(\d[\d\s]{2,5})",
        r"loyer[^\d]*(\d{3,4})\s*[€e](?:\s*/\s*mois)?",
        r"(\d{3,4})\s*[€e]\s*/\s*mois",
        r"rapport\s+locatif[^\d]*(\d[\d\s]{2,5})[€e]",
        r"revenus?\s+locatifs?[^\d]*(\d[\d\s]{2,5})[€e]",
    ]:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            try:
                v = int(re.sub(r"\s", "", m.group(1)))
                if 200 <= v <= 5000:
                    return v
            except ValueError:
                pass
    return None

def _regex_fallback_taxe(text):
    if not text: return None
    m = re.search(r"taxe\s+fonci[eè]re[^\d]*(\d[\d\s]*)\s*[€e]", text, re.IGNORECASE)
    if m:
        try: return int(re.sub(r"\s", "", m.group(1)))
        except ValueError: pass
    return None

def _regex_fallback_charges(text):
    if not text: return None
    for pat in [r"charges[^\d]*(\d{2,3})\s*[€e]\s*/\s*mois", r"charges\s+mensuelles[^\d]*(\d{2,3})"]:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            try: return int(re.sub(r"\s", "", m.group(1)))
            except ValueError: pass
    return None

def _regex_fallback_dpe(text):
    if not text: return None
    m = re.search(r"dpe\s*[:\-]?\s*([A-Ga-g])\b", text, re.IGNORECASE)
    return m.group(1).upper() if m else None

# ── Frais de notaire ───────────────────────────────────────────────
def frais_notaire(price: int) -> int:
    droits = price * 0.0580665
    if price <= 6500:       emol_ht = price * 0.03945
    elif price <= 17000:    emol_ht = 6500*0.03945 + (price-6500)*0.01627
    elif price <= 60000:    emol_ht = 6500*0.03945 + 10500*0.01627 + (price-17000)*0.01085
    else:                   emol_ht = 6500*0.03945 + 10500*0.01627 + 43000*0.01085 + (price-60000)*0.00814
    return int(droits + emol_ht*1.20 + price*0.001 + 1000)

# ── Rentabilité ────────────────────────────────────────────────────
def calc_yield_brut(loyer, prix_net):
    if None in (loyer, prix_net) or prix_net <= 0: return None
    return round(loyer * 12 / (prix_net + frais_notaire(prix_net)) * 100, 2)

def calc_yield_net(loyer, prix_net, taxe, charges):
    if None in (loyer, prix_net, taxe, charges): return None
    fn = frais_notaire(prix_net)
    total = prix_net + fn
    loyer_annuel = loyer * 12
    revenus_nets = loyer_annuel - taxe - charges*12*0.25 - loyer_annuel*0.5/12
    return round(revenus_nets / total * 100, 2)

def loyer_min_pour_rendement(prix_net: int) -> int:
    if not prix_net: return None
    total = prix_net + frais_notaire(prix_net)
    return math.ceil(total * MIN_YIELD / 100 / 12)

# ── Extraction des liens (vraies annonces uniquement) ──────────────
async def get_ad_links_from_page(page) -> list:
    return await page.evaluate("""() => {
        const seen = new Set();
        const out = [];
        for (const a of document.querySelectorAll('a[href]')) {
            const h = a.href;
            // Uniquement les vraies annonces avec ID numérique >= 7 chiffres
            if (/\\/ad\\/[^/]+\\/\\d{7,}/.test(h) && !seen.has(h)) {
                seen.add(h);
                out.push(h);
            }
        }
        return out;
    }""")

# ── Scraping page de résultats ─────────────────────────────────────
async def scrape_results_page(page) -> list:
    await page.wait_for_load_state("networkidle")
    await page.wait_for_timeout(human_delay(800, 1500))
    await human_scroll(page, "down")
    await page.wait_for_timeout(human_delay(400, 800))
    await human_scroll(page, "down")

    all_links, page_num = [], 1
    while True:
        await human_mouse_move(page)
        links = await get_ad_links_from_page(page)
        new = [l for l in links if l not in all_links]
        all_links.extend(new)
        print(f"  → Page {page_num} : {len(new)} lien(s) trouvé(s) (total {len(all_links)})")

        next_btn = await page.query_selector(
            "[data-qa-id='pagination_next_page'], a[aria-label='Page suivante'], "
            "a[rel='next'], [data-test-id='pagination-next'], button[aria-label*='suivant']"
        )
        if not next_btn: break
        prev_url = page.url
        try:
            await next_btn.scroll_into_view_if_needed()
            await page.wait_for_timeout(human_delay(400, 900))
            await next_btn.click()
            await page.wait_for_load_state("networkidle")
            await page.wait_for_timeout(human_delay(1500, 3000))
        except Exception as e:
            print(f"  [WARN] Pagination: {e}"); break
        if page.url == prev_url: break
        await detect_and_handle_block(page)
        page_num += 1

    return [{"url": l, "title": BLANK, "price": None, "location": BLANK,
             "surface": None, "loyer": None, "taxe_fonciere": None,
             "charges": None, "nb_pieces": None, "description": None} for l in all_links]

# ── Scraping détail annonce ────────────────────────────────────────
async def scrape_listing_detail(context, url: str) -> dict:
    data = {
        "url": url, "title": BLANK, "price": None, "location": BLANK,
        "surface": None, "loyer": None, "taxe_fonciere": None,
        "charges": None, "nb_pieces": None, "description": BLANK,
        "deja_loue": False, "bail_type": "inconnu",
        "loyer_potentiel": None, "dpe": None, "dpe_valeur": None,
    }
    page = await context.new_page()
    try:
        await page.wait_for_timeout(human_delay(700, 1800))
        await human_mouse_move(page)
        await page.goto(url, wait_until="domcontentloaded", timeout=30000)
        await page.wait_for_timeout(human_delay(1000, 2500))

        blocked = await detect_and_handle_block(page)
        if blocked:
            await page.goto(url, wait_until="domcontentloaded", timeout=30000)
            await page.wait_for_timeout(human_delay(1500, 3000))

        await human_scroll(page, "down")
        await page.wait_for_timeout(human_delay(500, 1200))
        await human_scroll(page, "down")

        # Titre
        for sel in ["h1","[data-qa-id='ad_title']","[data-test-id='ad-title']"]:
            el = await page.query_selector(sel)
            if el: data["title"] = _text(await el.inner_text()); break

        # Prix
        for sel in ["[data-qa-id='adview_price']","[data-test-id='ad-price']",
                    "[class*='price']","span[class*='Price']"]:
            el = await page.query_selector(sel)
            if el:
                data["price"] = parse_price(_text(await el.inner_text()))
                if data["price"]: break

        # Localisation
        for sel in ["[data-qa-id='adview_location_informations']","[data-test-id='ad-location']",
                    "[class*='location']","[class*='Location']"]:
            el = await page.query_selector(sel)
            if el: data["location"] = _text(await el.inner_text()).split("\n")[0]; break

        # Description
        for sel in ["[data-qa-id='adview_description_container']","[data-test-id='ad-description']",
                    "[class*='description']","[class*='Description']","div[itemprop='description']"]:
            el = await page.query_selector(sel)
            if el: data["description"] = _text(await el.inner_text()); break

        # Attributs structurés (surface, pièces, DPE, taxe)
        attr_els = await page.query_selector_all(
            "[data-qa-id='criteria_item'],[data-test-id='criteria-item'],"
            "[class*='criteria'],[class*='attribute'],[class*='Criteria']"
        )
        for attr_el in attr_els:
            text = _text(await attr_el.inner_text())
            tl = text.lower()
            if "surface" in tl or "m²" in tl:
                s = parse_surface(text)
                if s: data["surface"] = s
            elif "pièce" in tl or "piece" in tl:
                m = re.search(r"(\d+)", text)
                if m: data["nb_pieces"] = int(m.group(1))
            elif any(k in tl for k in ["dpe", "classe énergie", "consommation énergétique", "classe energie"]):
                m = re.search(r"\b([A-Ga-g])\b", text)
                if m: data["dpe"] = m.group(1).upper()
                m2 = re.search(r"(\d+)\s*kwh", text, re.IGNORECASE)
                if m2: data["dpe_valeur"] = int(m2.group(1))
            elif "taxe foncière" in tl or "taxe fonc" in tl:
                m = re.search(r"(\d[\d\s]*)\s*[€e]", text)
                if m:
                    try: data["taxe_fonciere"] = int(re.sub(r"\s", "", m.group(1)))
                    except ValueError: pass

        # DPE depuis badge dédié si pas encore trouvé
        if not data["dpe"]:
            for sel in ["[data-qa-id='energy_rate']","[data-test-id='dpe']",
                        "[class*='energy']","[class*='dpe']","[class*='DPE']"]:
                el = await page.query_selector(sel)
                if el:
                    txt = _text(await el.inner_text())
                    m = re.search(r"\b([A-Ga-g])\b", txt)
                    if m: data["dpe"] = m.group(1).upper(); break

        full_text = data["description"] or ""

        # ── Extraction IA ────────────────────────────────────────
        ai = await ai_extract_financials(full_text, data.get("price"),
                                         data.get("surface"), data.get("location"))
        if ai:
            loyer_hc   = ai.get("loyer_hc")
            loyer_cc   = ai.get("loyer_cc")
            charges_ai = ai.get("charges_mois")
            loyer_pot  = ai.get("loyer_potentiel")

            if loyer_hc:
                data["loyer"] = loyer_hc
            elif loyer_cc and charges_ai:
                data["loyer"] = loyer_cc - charges_ai
            elif loyer_cc:
                data["loyer"] = loyer_cc

            if data["loyer"] is None and loyer_pot:
                data["loyer"] = loyer_pot
                ai["loyer_source"] = "potentiel"

            if ai.get("taxe_fonciere") and not data["taxe_fonciere"]:
                data["taxe_fonciere"] = ai["taxe_fonciere"]
            if charges_ai:
                data["charges"] = charges_ai
            if ai.get("dpe") and not data["dpe"]:
                data["dpe"] = ai.get("dpe")
            if ai.get("dpe_valeur_kwh") and not data["dpe_valeur"]:
                data["dpe_valeur"] = ai.get("dpe_valeur_kwh")

            data["nb_lots"]       = ai.get("nb_lots", 1) or 1
            data["loyer_detail"]  = ai.get("loyer_detail")
            data["loyer_source"]  = ai.get("loyer_source", "inconnu")
            data["ai_notes"]      = ai.get("notes", "")
            data["deja_loue"]     = ai.get("deja_loue", False)
            data["bail_type"]     = ai.get("bail_type", "inconnu")
            data["loyer_potentiel"] = loyer_pot

            if data["loyer"]:
                src = data["loyer_source"]
                print(f"    [AI] loyer={data['loyer']}€/mois ({src})"
                      + (f" | DPE {data['dpe']}" if data["dpe"] else "")
                      + (f" | {data['nb_lots']} lots" if data["nb_lots"] > 1 else "")
                      + (f" | {str(data['ai_notes'])[:60]}" if data.get("ai_notes") else ""))
            else:
                print(f"    [AI] Loyer non trouvé (null)"
                      + (f" | {str(data.get('ai_notes',''))[:60]}" if data.get("ai_notes") else ""))
        else:
            # Fallback regex
            if data["loyer"] is None:
                data["loyer"] = _regex_fallback_loyer(full_text)
            if data["taxe_fonciere"] is None:
                data["taxe_fonciere"] = _regex_fallback_taxe(full_text)
            if data["charges"] is None:
                data["charges"] = _regex_fallback_charges(full_text)
            if data["dpe"] is None:
                data["dpe"] = _regex_fallback_dpe(full_text)
            data["nb_lots"] = 1; data["loyer_detail"] = None
            data["loyer_source"] = "regex"; data["ai_notes"] = ""

        # Estimation taxe foncière si complètement absente (indicatif)
        if data["taxe_fonciere"] is None and data.get("price"):
            data["taxe_fonciere_estimee"] = int(data["price"] * 0.008)
        else:
            data["taxe_fonciere_estimee"] = None

    except Exception as e:
        print(f"    [WARN] {url[:60]}: {e}")
    finally:
        await page.close()
    return data

# ── Enrichissement SANS filtre ─────────────────────────────────────
def enrich(listing: dict):
    """Retourne toujours un dict — aucun filtre sur le rendement ou le prix."""
    price  = listing.get("price")
    loyer  = listing.get("loyer")

    fn         = frais_notaire(price) if price else None
    total_cost = (price + fn) if price and fn else None
    yield_brut = calc_yield_brut(loyer, price)
    taxe       = listing.get("taxe_fonciere")
    taxe_est   = listing.get("taxe_fonciere_estimee")
    charges    = listing.get("charges")
    yield_net  = calc_yield_net(loyer, price, taxe, charges)
    surface    = listing.get("surface")
    price_m2   = round(price/surface, 0) if price and surface else None
    loyer_m2   = round(loyer/surface, 2) if loyer and surface else None
    loyer_min  = loyer_min_pour_rendement(price) if price else None

    nb_lots      = listing.get("nb_lots", 1) or 1
    loyer_detail = listing.get("loyer_detail")
    loyer_source = listing.get("loyer_source", BLANK)
    ai_notes     = listing.get("ai_notes", "")
    deja_loue    = listing.get("deja_loue", False)
    bail_type    = listing.get("bail_type", "inconnu")
    dpe          = listing.get("dpe") or BLANK
    dpe_valeur   = listing.get("dpe_valeur") or BLANK

    source_labels = {
        "hc": "HC extrait", "cc_moins_charges": "CC − charges",
        "cc": "CC (charges incluses)", "calcule": "Calculé",
        "potentiel": "Potentiel estimé", "regex": "Regex", "inconnu": BLANK,
    }
    loyer_label = source_labels.get(loyer_source, loyer_source)
    lots_str = BLANK
    if loyer_detail and isinstance(loyer_detail, list) and len(loyer_detail) > 1:
        lots_str = " + ".join(f"{v}€" for v in loyer_detail) + f" = {loyer}€"

    # Taxe : valeur réelle ou estimation
    taxe_affichee  = taxe if taxe else BLANK
    taxe_est_label = BLANK
    if taxe is None and taxe_est:
        taxe_est_label = f"≈{taxe_est}€ (estimé 0.8%)"

    return {
        "url":              listing.get("url", BLANK),
        "titre":            listing.get("title", BLANK),
        "ville":            listing.get("location", BLANK),
        "nb_pieces":        listing.get("nb_pieces") or BLANK,
        "surface_m2":       surface or BLANK,
        "dpe":              dpe,
        "dpe_kwh":          dpe_valeur,
        "prix_net":         price or BLANK,
        "frais_notaire":    fn or BLANK,
        "cout_total":       total_cost or BLANK,
        "prix_m2":          price_m2 or BLANK,
        "nb_lots":          nb_lots if nb_lots > 1 else BLANK,
        "loyer_hc":         loyer or BLANK,
        "loyer_min_requis": loyer_min or BLANK,
        "loyer_detail":     lots_str,
        "loyer_source":     loyer_label,
        "loyer_m2":         loyer_m2 or BLANK,
        "taxe_fonciere":    taxe_affichee,
        "taxe_fonciere_est":taxe_est_label,
        "charges_mois":     charges or BLANK,
        "rendement_brut":   yield_brut or BLANK,
        "rendement_net":    yield_net or BLANK,
        "deja_loue":        "Oui" if deja_loue else "Non",
        "bail_type":        bail_type if bail_type != "inconnu" else BLANK,
        "ai_notes":         (ai_notes[:120] if ai_notes else BLANK),
        "description":      (listing.get("description") or BLANK)[:300],
        "nouveau":          False,
        "date_scraping":    datetime.now().strftime("%Y-%m-%d %H:%M"),
    }

# ── Export HTML v3 ─────────────────────────────────────────────────
_HTML_STYLE = """<style>
*{box-sizing:border-box}
body{font-family:Arial,sans-serif;font-size:13px;margin:16px;background:#f0f0f0}
h1{color:#333;margin-bottom:8px}
.stats{background:#fff;padding:10px 18px;border-radius:8px;margin-bottom:12px;
       box-shadow:0 1px 4px rgba(0,0,0,.12);display:inline-block;font-size:13px}
.filters{background:#fff;padding:12px 16px;border-radius:8px;margin-bottom:12px;
         box-shadow:0 1px 4px rgba(0,0,0,.12);display:flex;flex-wrap:wrap;gap:8px;align-items:center}
.filters label{font-size:12px;font-weight:600;color:#555}
.filters input[type=text],.filters input[type=number]{
  padding:5px 9px;border:1px solid #ccc;border-radius:4px;width:130px}
.filters select{padding:5px 9px;border:1px solid #ccc;border-radius:4px}
.btn{padding:6px 14px;border:none;border-radius:4px;cursor:pointer;
     background:#e63946;color:#fff;font-weight:bold}
.btn:hover{background:#c1121f}
.table-wrap{overflow-x:auto}
table{border-collapse:collapse;width:100%;background:#fff;min-width:1400px;
      box-shadow:0 1px 4px rgba(0,0,0,.12);border-radius:8px;overflow:hidden}
th{background:#e63946;color:#fff;padding:8px 10px;cursor:pointer;
   white-space:nowrap;user-select:none;position:relative}
th:hover{background:#c1121f}
th.sort-asc::after{content:' ▲';font-size:10px}
th.sort-desc::after{content:' ▼';font-size:10px}
td{padding:6px 10px;border-bottom:1px solid #eee;vertical-align:top;white-space:nowrap}
td.wrap{white-space:normal;max-width:260px;font-size:12px;color:#555}
tr:hover td{background:#fff8f8}
tr.nouveau td{background:#fffde7}
.badge-new{background:#ff9800;color:#fff;padding:1px 6px;border-radius:8px;font-size:11px;font-weight:bold;margin-left:4px}
.badge-loue{background:#2d6a4f;color:#fff;padding:1px 6px;border-radius:8px;font-size:11px;font-weight:bold;margin-left:4px}
.rdt-top{color:#1a5c35;font-weight:bold;font-size:14px}
.rdt-high{color:#2d6a4f;font-weight:bold}
.rdt-med{color:#e07800;font-weight:bold}
.rdt-low{color:#999}
.dpe-badge{display:inline-block;padding:2px 9px;border-radius:4px;color:#fff;font-weight:bold;font-size:12px}
.dpe-fg{padding:2px 7px;border-radius:4px;color:#fff;font-weight:bold;font-size:12px;outline:2px solid red}
a{color:#e63946;text-decoration:none}a:hover{text-decoration:underline}
.est{color:#999;font-size:11px}
</style>"""

_HTML_SCRIPT = """<script>
let sCol=-1,sDir=1;
function sortTable(c,th){
  const tb=document.getElementById('tbl');
  const rows=Array.from(tb.tBodies[0].rows);
  if(sCol===c)sDir*=-1;else{sCol=c;sDir=1;}
  document.querySelectorAll('th').forEach(t=>t.classList.remove('sort-asc','sort-desc'));
  th.classList.add(sDir===1?'sort-asc':'sort-desc');
  rows.sort((a,b)=>{
    let va=a.cells[c].dataset.val??a.cells[c].innerText.replace(/[^0-9.\-]/g,'');
    let vb=b.cells[c].dataset.val??b.cells[c].innerText.replace(/[^0-9.\-]/g,'');
    if(va==='—'||va==='')return 1;
    if(vb==='—'||vb==='')return -1;
    const na=parseFloat(va),nb=parseFloat(vb);
    if(!isNaN(na)&&!isNaN(nb))return sDir*(na-nb);
    return sDir*va.localeCompare(vb,'fr');
  });
  rows.forEach(r=>tb.tBodies[0].appendChild(r));
}
function filterTable(){
  const q=document.getElementById('fSearch').value.toLowerCase();
  const minY=parseFloat(document.getElementById('fMinY').value)||0;
  const maxP=parseFloat(document.getElementById('fMaxP').value)||Infinity;
  const minP=parseFloat(document.getElementById('fMinP').value)||0;
  const dpeF=document.getElementById('fDPE').value;
  const onlyLoue=document.getElementById('fLoue').checked;
  const onlyNew=document.getElementById('fNew').checked;
  for(const r of document.getElementById('tbl').tBodies[0].rows){
    const yb=parseFloat(r.cells[20].dataset.val)||0;
    const pn=parseFloat(r.cells[7].dataset.val)||0;
    const dpe=r.cells[5].dataset.val||'';
    const loue=r.cells[22].innerText.trim()==='Oui';
    const isNew=r.classList.contains('nouveau');
    let ok=r.innerText.toLowerCase().includes(q)&&yb>=minY&&pn<=maxP&&pn>=minP;
    if(dpeF)ok=ok&&dpe===dpeF;
    if(onlyLoue)ok=ok&&loue;
    if(onlyNew)ok=ok&&isNew;
    r.style.display=ok?'':'none';
  }
}
// Tri par défaut : rendement brut décroissant au chargement
window.addEventListener('DOMContentLoaded',()=>{
  const th=document.querySelectorAll('th')[20];
  if(th){sCol=20;sDir=-1;sortTable(20,th);}
});
</script>"""

def _fmt(val, suf="", dec=None):
    if val == BLANK or val is None: return f'<span style="color:#bbb">{BLANK}</span>'
    if dec is not None: return f"{val:.{dec}f}{suf}"
    if isinstance(val, int): return f"{val:,}{suf}".replace(",", "\u202f")
    return f"{val}{suf}"

def _rdt_class(y):
    if y == BLANK or y is None: return "rdt-low"
    yf = float(y)
    if yf >= 15: return "rdt-top"
    if yf >= 10: return "rdt-high"
    if yf >= 8:  return "rdt-med"
    return "rdt-low"

def _dpe_badge(dpe):
    if dpe == BLANK or not dpe:
        return f'<span style="color:#bbb">{BLANK}</span>'
    color = DPE_COLORS.get(dpe.upper(), "#999")
    cls = "dpe-fg" if dpe.upper() in ("F", "G") else "dpe-badge"
    return f'<span class="{cls}" style="background:{color}" data-val="{dpe}">{dpe}</span>'

def export_html(listings: list, path: Path) -> None:
    headers = [
        "#","Titre","Ville","Pièces","Surface",
        "DPE","kWh/m²",
        "Prix net","Frais notaire","Coût total","Prix/m²",
        "Lots","Loyer HC","Loyer min 10%","Détail loyers","Source loyer","Loyer/m²",
        "Taxe foncière","Taxe (estimée)","Charges/mois",
        "Rdt brut","Rdt net",
        "Déjà loué","Type bail",
        "Notes IA","Description","Date scraping",
    ]
    ths = "".join(
        f'<th onclick="sortTable({i},this)">{h}</th>'
        for i, h in enumerate(headers)
    )

    rows = []
    for i, l in enumerate(listings, 1):
        badge_new  = '<span class="badge-new">NEW</span>'  if l.get("nouveau")         else ""
        badge_loue = '<span class="badge-loue">LOUÉ</span>' if l.get("deja_loue")=="Oui" else ""
        tr_cls = " class='nouveau'" if l.get("nouveau") else ""

        def c(val, suf="", dec=None, dv=None):
            dval = dv if dv is not None else (val if val != BLANK else "")
            return f'<td data-val="{dval}">{_fmt(val,suf,dec)}</td>'

        yb  = l["rendement_brut"]
        yn  = l["rendement_net"]
        dpe = l.get("dpe", BLANK)

        src = l.get("loyer_source", BLANK)
        src_c = {"HC extrait":"#2d6a4f","CC − charges":"#1a6ba0","CC (charges incluses)":"#e07800",
                 "Potentiel estimé":"#9b59b6","Regex":"#888"}.get(src,"#555")
        src_html = f'<span style="color:{src_c};font-weight:bold;font-size:11px">{src}</span>' if src != BLANK else _fmt(BLANK)

        taxe_af  = l.get("taxe_fonciere", BLANK)
        taxe_est = l.get("taxe_fonciere_est", BLANK)

        rows.append(f"""<tr{tr_cls}>
<td>{i}</td>
<td><a href="{l['url']}" target="_blank">{str(l['titre'])[:45]}{badge_new}{badge_loue}</a></td>
<td>{l['ville']}</td>
{c(l['nb_pieces'])}{c(l['surface_m2'],' m²')}
<td data-val="{dpe if dpe!=BLANK else ''}">{_dpe_badge(dpe)}</td>
{c(l['dpe_kwh'],' kWh')}
{c(l['prix_net'],' €',dv=l['prix_net'] if l['prix_net']!=BLANK else '')}
{c(l['frais_notaire'],' €')}{c(l['cout_total'],' €')}{c(l['prix_m2'],' €/m²')}
{c(l['nb_lots'])}
{c(l['loyer_hc'],' €/mois')}
<td data-val="{l['loyer_min_requis'] if l['loyer_min_requis']!=BLANK else ''}" style="color:#9b59b6;font-weight:bold">{_fmt(l['loyer_min_requis'],' €/mois')}</td>
<td class="wrap">{l.get('loyer_detail',BLANK) if l.get('loyer_detail',BLANK)!=BLANK else _fmt(BLANK)}</td>
<td>{src_html}</td>
{c(l['loyer_m2'],' €/m²')}
{c(taxe_af,' €/an')}
<td class="est">{taxe_est if taxe_est!=BLANK else _fmt(BLANK)}</td>
{c(l['charges_mois'],' €/mois')}
<td data-val="{yb if yb!=BLANK else ''}"><span class="{_rdt_class(yb)}">{_fmt(yb,'%',2)}</span></td>
<td data-val="{yn if yn!=BLANK else ''}"><span class="{_rdt_class(yn)}">{_fmt(yn,'%',2)}</span></td>
<td>{l.get('deja_loue',BLANK)}</td>
<td>{l.get('bail_type',BLANK)}</td>
<td class="wrap">{l.get('ai_notes',BLANK) if l.get('ai_notes',BLANK)!=BLANK else _fmt(BLANK)}</td>
<td class="wrap">{str(l['description'])[:180]}</td>
<td>{l.get('date_scraping',BLANK)}</td>
</tr>""")

    total    = len(listings)
    bruts    = [float(l["rendement_brut"]) for l in listings if l["rendement_brut"] != BLANK]
    avg      = round(sum(bruts)/len(bruts), 2) if bruts else 0
    top      = max(bruts) if bruts else 0
    nouveaux = sum(1 for l in listings if l.get("nouveau"))
    loues    = sum(1 for l in listings if l.get("deja_loue") == "Oui")
    avec_rdt = sum(1 for l in listings if l["rendement_brut"] != BLANK)
    sur_10   = sum(1 for l in listings if l["rendement_brut"] != BLANK and float(l["rendement_brut"]) >= MIN_YIELD)

    html = f"""<!DOCTYPE html>
<html lang="fr"><head><meta charset="utf-8">
<title>LeBonCoin — Analyse immobilière</title>
{_HTML_STYLE}
</head><body>
<h1>LeBonCoin — Analyse immobilière complète</h1>
<div class="stats">
  <strong>{total}</strong> biens &nbsp;|&nbsp;
  Avec rendement calculé : <strong>{avec_rdt}</strong> &nbsp;|&nbsp;
  ≥ {MIN_YIELD}% brut : <strong style="color:#2d6a4f">{sur_10}</strong> &nbsp;|&nbsp;
  Moy. brut : <strong>{avg}%</strong> &nbsp;|&nbsp;
  Meilleur : <strong class="rdt-top">{top}%</strong> &nbsp;|&nbsp;
  Nouveaux : <strong>{nouveaux}</strong> &nbsp;|&nbsp;
  Loués : <strong>{loues}</strong> &nbsp;|&nbsp;
  {datetime.now().strftime("%d/%m/%Y %H:%M")}
</div>
<div class="filters">
  <label>Recherche <input type="text" id="fSearch" placeholder="ville, titre…" oninput="filterTable()"></label>
  <label>Rdt brut min % <input type="number" id="fMinY" placeholder="0" oninput="filterTable()" step="0.5"></label>
  <label>Prix min € <input type="number" id="fMinP" placeholder="0" oninput="filterTable()" step="1000"></label>
  <label>Prix max € <input type="number" id="fMaxP" placeholder="∞" oninput="filterTable()" step="1000"></label>
  <label>DPE <select id="fDPE" onchange="filterTable()">
    <option value="">Tous</option>
    <option>A</option><option>B</option><option>C</option>
    <option>D</option><option>E</option><option>F</option><option>G</option>
  </select></label>
  <label><input type="checkbox" id="fLoue" onchange="filterTable()"> Déjà loué</label>
  <label><input type="checkbox" id="fNew"  onchange="filterTable()"> Nouveaux seulement</label>
  <button class="btn" onclick="filterTable()">Filtrer</button>
</div>
<div class="table-wrap">
<table id="tbl">
<thead><tr>{ths}</tr></thead>
<tbody>{''.join(rows)}</tbody>
</table>
</div>
{_HTML_SCRIPT}
</body></html>"""
    path.write_text(html, encoding="utf-8")
    print(f"  HTML : {path}")

# ── Export Excel (.xlsx) ───────────────────────────────────────────
def export_xlsx(listings: list, path: Path) -> None:
    if not XLSX_OK:
        print("  [WARN] openpyxl non installé — export Excel ignoré. Faites : pip install openpyxl")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Résultats"

    headers = [
        "#","URL","Titre","Ville","Pièces","Surface m²","DPE","kWh/m²",
        "Prix net €","Frais notaire €","Coût total €","Prix/m² €",
        "Lots","Loyer HC €/mois","Loyer min 10% €/mois","Détail loyers","Source loyer","Loyer/m²",
        "Taxe foncière €/an","Taxe estimée","Charges €/mois",
        "Rdt brut %","Rdt net %",
        "Déjà loué","Type bail","Notes IA","Description","Date scraping",
    ]

    # Couleurs
    RED_FILL    = PatternFill("solid", fgColor="E63946")
    GREEN_FILL  = PatternFill("solid", fgColor="D4EDDA")
    ORANGE_FILL = PatternFill("solid", fgColor="FFF3CD")
    GREY_FILL   = PatternFill("solid", fgColor="F8F9FA")
    WHITE_FONT  = Font(bold=True, color="FFFFFF")
    BOLD        = Font(bold=True)
    CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # En-têtes
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = RED_FILL
        cell.font = WHITE_FONT
        cell.alignment = CENTER

    # Données
    for row_i, l in enumerate(listings, 2):
        yb = l["rendement_brut"]
        yb_val = float(yb) if yb != BLANK else None

        def val(key):
            v = l.get(key, BLANK)
            return None if v == BLANK else v

        row_data = [
            row_i - 1,
            l["url"],
            str(l["titre"])[:80],
            l["ville"],
            val("nb_pieces"),
            val("surface_m2"),
            val("dpe"),
            val("dpe_kwh"),
            val("prix_net"),
            val("frais_notaire"),
            val("cout_total"),
            val("prix_m2"),
            val("nb_lots"),
            val("loyer_hc"),
            val("loyer_min_requis"),
            val("loyer_detail"),
            val("loyer_source"),
            val("loyer_m2"),
            val("taxe_fonciere"),
            val("taxe_fonciere_est"),
            val("charges_mois"),
            yb_val,
            (float(l["rendement_net"]) if l["rendement_net"] != BLANK else None),
            l.get("deja_loue", "Non"),
            val("bail_type"),
            str(val("ai_notes") or "")[:120],
            str(val("description") or "")[:200],
            l.get("date_scraping", ""),
        ]

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_i, column=col, value=value)
            # Alterner fond de ligne
            if row_i % 2 == 0:
                cell.fill = GREY_FILL
            # URL en hyperlien
            if col == 2 and value:
                cell.hyperlink = value
                cell.value = "Voir annonce"
                cell.font = Font(color="E63946", underline="single")

        # Mise en forme conditionnelle manuelle sur rendement brut (colonne 22)
        rdt_cell = ws.cell(row=row_i, column=22)
        if yb_val is not None:
            if yb_val >= MIN_YIELD:
                rdt_cell.fill = GREEN_FILL
                rdt_cell.font = Font(bold=True, color="1A5C35")
            elif yb_val >= 8:
                rdt_cell.fill = ORANGE_FILL
                rdt_cell.font = Font(bold=True, color="856404")
            rdt_cell.number_format = '0.00"%"'

        # DPE rouge si F ou G (colonne 7)
        dpe_val = l.get("dpe", BLANK)
        if dpe_val in ("F", "G"):
            dpe_cell = ws.cell(row=row_i, column=7)
            dpe_cell.font = Font(bold=True, color="CC0000")

    # AutoFilter sur toute la plage
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # Freeze la ligne d'en-tête
    ws.freeze_panes = "A2"

    # Largeurs de colonnes approximatives
    col_widths = [4,14,35,20,7,10,6,8,13,14,13,11,6,15,16,18,14,10,15,18,12,10,10,10,12,25,30,18]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # Hauteur de la ligne d'en-tête
    ws.row_dimensions[1].height = 35

    wb.save(path)
    print(f"  XLSX : {path}")

# ── Export CSV (backup) ────────────────────────────────────────────
def export_csv(listings: list, path: Path) -> None:
    if not listings: return
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(listings[0].keys()))
        writer.writeheader()
        writer.writerows(listings)
    print(f"  CSV  : {path}")

# ── WhatsApp ───────────────────────────────────────────────────────
def send_whatsapp(message: str) -> None:
    if not WA_PHONE or not WA_APIKEY: return
    try:
        encoded = urllib.parse.quote(message)
        url = f"https://api.callmebot.com/whatsapp.php?phone={WA_PHONE}&text={encoded}&apikey={WA_APIKEY}"
        with urllib.request.urlopen(url, timeout=10) as resp:
            print(f"  WhatsApp : {'OK' if resp.status==200 else f'HTTP {resp.status}'}")
    except Exception as e:
        print(f"  WhatsApp : erreur ({e})")

# ── Main ───────────────────────────────────────────────────────────
async def main():
    print("=" * 62)
    print(" LeBonCoin Analyzer v3 — Tous résultats + Cache + DPE")
    print("=" * 62)
    ai_label = {"openai":"OpenAI gpt-4o-mini","anthropic":"Anthropic Haiku"}.get(AI_PROVIDER,"AUCUN (regex)")
    print(f" IA      : {ai_label}")
    print(f" Prix max: {MAX_PRICE:,} € | Cible rendement : {MIN_YIELD}%")
    print(f" CDP     : {CDP_URL}")
    if not XLSX_OK:
        print(" [INFO] openpyxl absent — pas d'export Excel. Installez : pip install openpyxl")
    print()

    seen  = load_seen()
    cache = load_cache()
    print(f" Cache : {len(cache)} annonce(s) déjà scrapée(s) — seront ignorées")
    print()

    async with async_playwright() as p:
        try:
            browser = await p.chromium.connect_over_cdp(CDP_URL)
        except Exception as e:
            print(f"ERREUR : impossible de se connecter a Chrome ({e})")
            print("  → Lancez lancer_chrome.bat d'abord")
            sys.exit(1)

        contexts = browser.contexts
        if not contexts:
            print("ERREUR : aucun contexte navigateur."); await browser.close(); sys.exit(1)

        context = contexts[0]
        await load_cookies(context)
        await apply_stealth(context, random.choice(_USER_AGENTS), random.choice(_VIEWPORTS))

        pages = context.pages
        page  = pages[0] if pages else await context.new_page()
        await detect_and_handle_block(page)

        page_type = await detect_page_type(page)
        print(f" Page détectée : {page_type} ({page.url[:80]})\n")

        result_urls = []
        if page_type == "saved_searches":
            print(" Récupération des recherches sauvegardées...")
            result_urls = await get_saved_search_urls(page)
            if not result_urls:
                print(" Aucune trouvée. Naviguez manuellement puis appuyez sur Entrée...")
                input(); result_urls = [page.url]
            else:
                print(f" {len(result_urls)} recherche(s) :")
                for u in result_urls: print(f"   * {u[:80]}")
        elif page_type == "results":
            result_urls = [page.url]
            print(f" Page de résultats : {page.url[:80]}")
        else:
            print(" Page non reconnue. Naviguez puis appuyez sur Entrée...")
            input(); result_urls = [page.url]
        print()

        # Scraping des listes
        all_raw = []
        if page_type == "results" and result_urls and page.url == result_urls[0]:
            print("Extraction sans navigation...")
            all_raw.extend(await scrape_results_page(page))
        else:
            for target_url in result_urls:
                print(f"Navigation vers : {target_url[:80]}")
                if page.url != target_url:
                    await page.goto(target_url, wait_until="networkidle")
                    await page.wait_for_timeout(human_delay(2000, 4000))
                    await detect_and_handle_block(page)
                all_raw.extend(await scrape_results_page(page))

        await save_cookies(context)

        # Dédupliquer
        seen_batch, unique_raw = {}, []
        for l in all_raw:
            if l and l.get("url") and l["url"] not in seen_batch:
                seen_batch[l["url"]] = True
                unique_raw.append(l)

        cached_count = sum(1 for l in unique_raw if l.get("url","") in cache)
        new_count    = len(unique_raw) - cached_count
        print(f"\n{len(unique_raw)} annonce(s) uniques — {cached_count} en cache (skip) — {new_count} à scraper\n")

        enriched     = []
        new_listings = []
        scrape_i     = 0

        for i, card in enumerate(unique_raw, 1):
            url = card.get("url", "")

            if url in cache:
                # Utiliser les données en cache — pas de scraping
                print(f"  [{i}/{len(unique_raw)}] [CACHE] {url[:65]}")
                detail = cache[url]
            else:
                scrape_i += 1
                print(f"  [{i}/{len(unique_raw)}] {url[:65]}")
                # Pauses humaines seulement sur les nouvelles annonces
                if scrape_i > 1 and scrape_i % 8 == 0:
                    pause = random.randint(4000, 9000)
                    print(f"    [pause {pause//1000}s]")
                    await page.wait_for_timeout(pause)
                if scrape_i > 1 and scrape_i % 25 == 0:
                    pause = random.randint(15000, 30000)
                    print(f"    [pause longue {pause//1000}s]")
                    await page.wait_for_timeout(pause)
                detail = await scrape_listing_detail(context, url)
                cache[url] = detail
                save_cache(cache)

            merged = {**card, **{k: v for k, v in detail.items() if v is not None and v != BLANK}}
            result = enrich(merged)

            is_new = url not in seen
            result["nouveau"] = is_new
            if is_new:
                new_listings.append(result)
                seen.add(url)

            enriched.append(result)

            # Affichage synthétique
            yb  = result["rendement_brut"]
            dpe = result.get("dpe", BLANK)
            dpe_str = f" | DPE {dpe}" if dpe != BLANK else ""
            if yb != BLANK:
                print(f"    ✓ {result.get('prix_net',BLANK):,} € | loyer {result['loyer_hc']} €/mois | brut {yb}%{dpe_str}")
            else:
                lm = result.get("loyer_min_requis", BLANK)
                print(f"    ○ Pas de loyer trouvé | loyer min {MIN_YIELD}% : {lm} €/mois{dpe_str}")

        await save_cookies(context)

        # Tri par rendement brut décroissant (BLANK à la fin)
        enriched.sort(key=lambda x: float(x["rendement_brut"]) if x["rendement_brut"] != BLANK else -999, reverse=True)

        avec_rdt = sum(1 for l in enriched if l["rendement_brut"] != BLANK)
        sur_10   = sum(1 for l in enriched if l["rendement_brut"] != BLANK and float(l["rendement_brut"]) >= MIN_YIELD)

        print(f"\n{'='*62}")
        print(f" {len(enriched)} biens | {avec_rdt} avec rendement | {sur_10} ≥ {MIN_YIELD}% | {len(new_listings)} nouveau(x)")
        print(f"{'='*62}\n")

        ts        = datetime.now().strftime("%Y%m%d_%H%M%S")
        html_path = OUTPUT_DIR / f"rapport_{ts}.html"
        csv_path  = OUTPUT_DIR / f"rapport_{ts}.csv"
        xlsx_path = OUTPUT_DIR / f"rapport_{ts}.xlsx"

        export_html(enriched, html_path)
        export_xlsx(enriched, xlsx_path)
        export_csv(enriched, csv_path)

        if new_listings and WA_PHONE and WA_APIKEY:
            sur_10_new = [l for l in new_listings if l["rendement_brut"] != BLANK and float(l["rendement_brut"]) >= MIN_YIELD]
            if sur_10_new:
                lines = [f"LeBonCoin — {len(sur_10_new)} nouveau(x) bien(s) ≥ {MIN_YIELD}% brut :"]
                for l in sur_10_new[:5]:
                    dpe_str = f" DPE:{l.get('dpe','-')}" if l.get("dpe",BLANK) != BLANK else ""
                    lines.append(f"* {l['ville']} — {l['prix_net']:,}€ — {l['rendement_brut']}%{dpe_str}")
                if len(sur_10_new) > 5:
                    lines.append(f"...et {len(sur_10_new)-5} autre(s).")
                send_whatsapp("\n".join(lines))

        save_seen(seen)

        print(f"\nHTML  : {html_path.resolve()}")
        print(f"Excel : {xlsx_path.resolve()}")
        print(f"CSV   : {csv_path.resolve()}")
        print("Terminé.")
        await browser.close()


if __name__ == "__main__":
    asyncio.run(main())
